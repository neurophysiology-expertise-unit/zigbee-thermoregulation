"""Read session/recording `.jsonl` files and reduce them to numbers.

Headless and stdlib-only (matplotlib is imported lazily, only for `--plots`),
so it runs on the analysis machine, not just the rig: the GUI's Review tab and
this module read the same files through the same parser, which lives here.

    python -m mouse_thermo.recordings <file-or-dir>... [--csv out.csv] [--plots dir]

One row per session. What the rows are FOR: the control-quality numbers a
manuscript reports -- how fast the rig reached the setpoint, how tightly it
held it, how hard the actuator worked -- plus the body-sensor dropout split by
actuator state, which is the quantitative form of the EMI finding in CLAUDE.md
(the RFID reader goes blind while the lamp is on).

This module ONLY reads. It never writes into a recording, and nothing here is
in the control path.
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import sys
from typing import Optional

NAN = float("nan")


def _num(x) -> float:
    return float(x) if isinstance(x, (int, float)) else NAN


def _ok(v: float) -> bool:
    """True for a real number. NaN != NaN is the repo's idiom for "missing"."""
    return v == v


def parse(path: str) -> dict:
    """Read a session/recording .jsonl into column arrays.

    Tolerant of missing keys (older files) and of a truncated last line (a
    session killed mid-write) -- those are skipped, never fatal. A recording
    of an interrupted experiment is still evidence.
    """
    config: dict = {}
    t_mono, t_wall, body, ambient, lamp, power = [], [], [], [], [], []
    state, reason, body_sp, amb_sp = [], [], [], []
    body_age, rfid_age, lamp_state, ground_truth, record_mode = [], [], [], [], []
    raw_body, pulse = [], []
    events: list[dict] = []

    with open(path, "r") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue  # partial trailing line from an interrupted session
            typ = rec.get("type")
            if typ == "session_start":
                config = rec.get("config", {}) or {}
            elif typ == "event":
                events.append(rec)
            elif typ == "sample":
                t_mono.append(_num(rec.get("t_mono")))
                t_wall.append(_num(rec.get("t_wall")))
                body.append(_num(rec.get("body_c")))
                ambient.append(_num(rec.get("ambient_c")))
                lamp.append(bool(rec.get("lamp_cmd")))
                power.append(_num(rec.get("power_w")))
                state.append(rec.get("state") or "")
                reason.append(rec.get("reason") or "")
                body_sp.append(_num(rec.get("body_setpoint_c")))
                amb_sp.append(_num(rec.get("ambient_setpoint_c")))
                body_age.append(_num(rec.get("body_age_s")))
                rfid_age.append(_num(rec.get("raw_rfid_age_s")))
                # raw_rfid_c is the reader's own value, BEFORE the plausibility
                # gate. When body_valid_range excludes a real excursion, this
                # column is the only place the excursion survives.
                raw_body.append(_num(rec.get("raw_rfid_c")))
                pulse.append(bool(rec.get("pulse_active")))
                lamp_state.append(rec.get("lamp_state"))
                ground_truth.append(rec.get("ground_truth") or "")
                record_mode.append(rec.get("record_mode") or "")

    t0 = next((v for v in t_mono if _ok(v)), 0.0)
    t = [(v - t0) if _ok(v) else NAN for v in t_mono]
    return {
        "path": path, "config": config, "t": t,
        "body": body, "ambient": ambient, "lamp": lamp, "power": power,
        "state": state, "reason": reason,
        "body_sp": body_sp, "amb_sp": amb_sp,
        # extras beyond what the Review tab plots
        "body_age": body_age, "rfid_age": rfid_age, "lamp_state": lamp_state,
        "raw_body": raw_body, "pulse": pulse, "t_wall": t_wall,
        "ground_truth": ground_truth, "record_mode": record_mode,
        "events": events,
    }


def control_stats_text(rec: dict) -> str:
    """The human-readable per-session block the GUI's Review tab shows."""
    ctrl = (rec.get("config") or {}).get("control", {}) or {}
    mode = ctrl.get("mode", "heat")
    body_db = float(ctrl.get("body_deadband_c", 0.3) or 0.3)
    amb_db = float(ctrl.get("ambient_deadband_c", 0.5) or 0.5)

    def stat_line(name, arr, sp_arr, db):
        pairs = [(v, s) for v, s in zip(arr, sp_arr) if _ok(v)]
        vals = [v for v, _ in pairs]
        if not vals:
            return None
        n = len(vals)
        m = sum(vals) / n
        sd = (sum((v - m) ** 2 for v in vals) / n) ** 0.5
        withsp = [(v, s) for v, s in pairs if _ok(s)]
        in_band = (100.0 * sum(1 for v, s in withsp if abs(v - s) <= db) / len(withsp)
                   if withsp else None)
        band_txt = f", {in_band:.0f}% within ±{db:g}°C of setpoint" if in_band is not None else ""
        return (f"{name}: mean {m:.2f}°C, SD {sd:.2f}, range {min(vals):.2f}–{max(vals):.2f}"
                f"{band_txt}  (n={n})")

    lines = [f"Mode: {mode}"]
    for name, key, sp_key, db in (("Body", "body", "body_sp", body_db),
                                  ("Ambient", "ambient", "amb_sp", amb_db)):
        ln = stat_line(name, rec[key], rec[sp_key], db)
        if ln:
            lines.append(ln)

    lamp = rec["lamp"]
    t = [x for x in rec["t"] if _ok(x)]
    if lamp:
        duty = duty_pct(rec["t"], lamp) or 0.0
        dur = (t[-1] - t[0]) if len(t) >= 2 else 0.0
        lines.append(f"Actuator ON {duty:.0f}% of the time over {dur:.0f}s "
                     f"({len(lamp)} samples)")
    # A LOCKOUT anywhere is worth flagging explicitly.
    n_lockout = sum(1 for s in rec["state"] if s == "LOCKOUT")
    if n_lockout:
        lines.append(f"⚠ {n_lockout} sample(s) in LOCKOUT")
    return "\n".join(lines)


# --- metrics ----------------------------------------------------------------

def duty_pct(t: list[float], on: list[bool]) -> Optional[float]:
    """Fraction of TIME the actuator was commanded on, not fraction of samples.

    The loop wakes on pulse edges, so samples are not evenly spaced: a 5 s ON
    burst and a 1 s OFF gap produce one sample each, and counting samples would
    call that a 50% duty when it is 83%. Integrate over the intervals instead.
    """
    tot = held = 0.0
    for i in range(len(on) - 1):
        if not (_ok(t[i]) and _ok(t[i + 1])):
            continue
        dt = t[i + 1] - t[i]
        if dt <= 0:
            continue
        tot += dt
        if on[i]:
            held += dt
    if tot <= 0:
        # a single sample, or no usable timestamps: fall back to the sample
        # fraction and let the caller see it rather than reporting nothing
        return (100.0 * sum(1 for x in on if x) / len(on)) if on else None
    return 100.0 * held / tot


def _stats(vals: list[float]) -> dict:
    v = [x for x in vals if _ok(x)]
    if not v:
        return {"n": 0, "mean": None, "sd": None, "min": None, "max": None}
    n = len(v)
    m = sum(v) / n
    return {"n": n, "mean": m,
            "sd": (sum((x - m) ** 2 for x in v) / n) ** 0.5,
            "min": min(v), "max": max(v)}


def _median(vals: list[float]) -> Optional[float]:
    v = sorted(x for x in vals if _ok(x))
    if not v:
        return None
    mid = len(v) // 2
    return v[mid] if len(v) % 2 else (v[mid - 1] + v[mid]) / 2.0


def _pct_in_band(vals, sps, db) -> Optional[float]:
    pairs = [(v, s) for v, s in zip(vals, sps) if _ok(v) and _ok(s)]
    if not pairs:
        return None
    return 100.0 * sum(1 for v, s in pairs if abs(v - s) <= db) / len(pairs)


def _time_to_band(t, vals, sps, db) -> Optional[float]:
    """Seconds from the first sample until the source first sits in band.

    Reported as None when it never does, and 0.0 when it started there --
    a session that began at setpoint measures nothing about approach speed,
    so the two cases must not collapse into the same number.
    """
    for ti, v, s in zip(t, vals, sps):
        if _ok(ti) and _ok(v) and _ok(s) and abs(v - s) <= db:
            t_first = next((x for x in t if _ok(x)), None)
            return None if t_first is None else ti - t_first
    return None


def _modal(vals: list[str]) -> Optional[str]:
    seen = [v for v in vals if v]
    if not seen:
        return None
    return max(set(seen), key=seen.count)


def _episodes(states: list[str], name: str) -> int:
    """Entries INTO a state, not samples spent in it."""
    n, prev = 0, None
    for s in states:
        if s == name and prev != name:
            n += 1
        prev = s
    return n


def _dropout_pct(vals: list[float], mask: list[bool]) -> Optional[float]:
    sel = [v for v, m in zip(vals, mask) if m]
    if not sel:
        return None
    return 100.0 * sum(1 for v in sel if not _ok(v)) / len(sel)


def summarize(rec: dict) -> dict:
    """One flat row of numbers per session, for a CSV a manuscript can cite."""
    cfg = rec.get("config") or {}
    ctrl = cfg.get("control", {}) or {}
    body_db = float(ctrl.get("body_deadband_c", 0.3) or 0.3)
    amb_db = float(ctrl.get("ambient_deadband_c", 0.5) or 0.5)

    t = [x for x in rec["t"] if _ok(x)]
    lamp = rec["lamp"]
    on = [bool(x) for x in lamp]
    off = [not x for x in on]

    b, a = _stats(rec["body"]), _stats(rec["ambient"])
    rb = _stats(rec["raw_body"])
    row = {
        "session": os.path.basename(rec["path"]),
        "mode": ctrl.get("mode", "heat"),
        "record_mode": _modal(rec["record_mode"]),
        "ground_truth": _modal(rec["ground_truth"]),
        "duration_s": (t[-1] - t[0]) if len(t) >= 2 else 0.0,
        "n_samples": len(lamp),
        "body_setpoint_c": _median(rec["body_sp"]),
        "ambient_setpoint_c": _median(rec["amb_sp"]),
        "body_mean_c": b["mean"], "body_sd_c": b["sd"],
        "body_min_c": b["min"], "body_max_c": b["max"], "body_n": b["n"],
        "ambient_mean_c": a["mean"], "ambient_sd_c": a["sd"],
        "ambient_min_c": a["min"], "ambient_max_c": a["max"], "ambient_n": a["n"],
        "body_pct_in_band": _pct_in_band(rec["body"], rec["body_sp"], body_db),
        "ambient_pct_in_band": _pct_in_band(rec["ambient"], rec["amb_sp"], amb_db),
        "body_time_to_band_s": _time_to_band(rec["t"], rec["body"], rec["body_sp"], body_db),
        "ambient_time_to_band_s": _time_to_band(rec["t"], rec["ambient"], rec["amb_sp"], amb_db),
        "actuator_duty_pct": duty_pct(rec["t"], on),
        "actuator_transitions": sum(1 for i in range(1, len(on)) if on[i] != on[i - 1]),
        "lockout_samples": sum(1 for s in rec["state"] if s == "LOCKOUT"),
        "lockout_episodes": _episodes(rec["state"], "LOCKOUT"),
        "fallback_episodes": _episodes(rec["state"], "FALLBACK"),
        # The EMI finding, as a number: the body reader goes blind under the
        # actuator. Compare the two dropout columns -- a large ON/OFF gap is
        # the reader being jammed, not an animal moving off the probe.
        "body_dropout_pct": _dropout_pct(rec["body"], [True] * len(rec["body"])),
        "body_dropout_pct_actuator_on": _dropout_pct(rec["body"], on),
        "body_dropout_pct_actuator_off": _dropout_pct(rec["body"], off),
        "rfid_age_median_s_actuator_on": _median([v for v, m in zip(rec["rfid_age"], on) if m]),
        "rfid_age_median_s_actuator_off": _median([v for v, m in zip(rec["rfid_age"], off) if m]),
        # The reader's own values, pre-gate. body_gated_pct is the share of
        # samples where the reader HAD a value and the plausibility gate
        # rejected it -- a high number means body_valid_range is excluding the
        # very excursion being studied, not that the reader failed.
        "raw_body_first_c": next((v for v in rec["raw_body"] if _ok(v)), None),
        "raw_body_last_c": next((v for v in reversed(rec["raw_body"]) if _ok(v)), None),
        "raw_body_min_c": rb["min"], "raw_body_max_c": rb["max"],
        "body_gated_pct": (100.0 * sum(1 for v, b in zip(rec["raw_body"], rec["body"])
                                       if _ok(v) and not _ok(b)) / len(rec["raw_body"]))
                          if rec["raw_body"] else None,
        "n_events": len(rec["events"]),
    }
    return row


def write_csv(rows: list[dict], path: str) -> None:
    if not rows:
        raise SystemExit("no sessions parsed; nothing to write")
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


XLSX_README = [
    "One sheet per session, plus Summary (one row per session).",
    "",
    "chip = the implanted temperature transponder, read by the RFID reader.",
    "chip_temp_raw_c  what the reader reported (use this one).",
    "chip_temp_c      the same value AFTER the software plausibility gate; empty when",
    "                 the reading fell outside sensors.body_valid_range in the config.",
    "chip_age_s       seconds since the reader last produced a value.",
    "ambient_c        box air temperature (ESP32 probe).",
    "actuator_on      1 = the smart plug was commanded ON (heat lamp or Peltier).",
    "state            controller state: NORMAL / FALLBACK / LOCKOUT.",
    "time_utc         wall clock in UTC; time_s counts from the session's first sample.",
    "",
    "Samples are not evenly spaced: the loop also wakes on actuator pulse edges.",
    "Summary duty is time-weighted for that reason.",
    "Generated by mouse_thermo.recordings from the raw .jsonl; the .jsonl is the record.",
]

_SHEET_COLS = [
    ("time_s", "t"), ("time_utc", "t_wall"),
    ("chip_temp_raw_c", "raw_body"), ("chip_temp_c", "body"), ("chip_age_s", "rfid_age"),
    ("ambient_c", "ambient"), ("actuator_on", "lamp"), ("pulse_active", "pulse"),
    ("state", "state"), ("reason", "reason"),
    ("chip_setpoint_c", "body_sp"), ("ambient_setpoint_c", "amb_sp"),
    ("record_mode", "record_mode"), ("ground_truth", "ground_truth"),
]


def write_xlsx(recs: list[dict], path: str) -> None:
    """Excel workbook for people who will not open a .jsonl. openpyxl is
    imported here so the rest of the module stays stdlib-only."""
    import datetime
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    def cell(key, v):
        if isinstance(v, float) and not _ok(v):
            return None                      # empty cell, never a fake 0
        if key == "t_wall" and v is not None:
            return datetime.datetime.fromtimestamp(v, datetime.timezone.utc).replace(tzinfo=None)
        if isinstance(v, bool):
            return int(v)
        return v

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in XLSX_README:
        ws.append([line])
    ws.column_dimensions["A"].width = 90

    rows = [summarize(r) for r in recs]
    ws = wb.create_sheet("Summary")
    ws.append(list(rows[0].keys()))
    for row in rows:
        ws.append(list(row.values()))

    used = set()
    for rec in recs:
        name = os.path.splitext(os.path.basename(rec["path"]))[0][:31]
        while name in used:                  # Excel: unique, <= 31 chars
            name = name[:28] + f"_{len(used)}"
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([c for c, _ in _SHEET_COLS])
        for i in range(len(rec["t"])):
            ws.append([cell(k, rec[k][i]) for _, k in _SHEET_COLS])
        for j, (col, key) in enumerate(_SHEET_COLS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = 20 if key == "t_wall" else max(11, len(col) + 2)
            if key == "t_wall":
                for c in ws[get_column_letter(j)][1:]:
                    c.number_format = "yyyy-mm-dd hh:mm:ss"
        ws.freeze_panes = "A2"
    wb.save(path)


def plot(rec: dict, out_path: str) -> None:
    """One overview figure per session. matplotlib is imported here, not at
    module scope, so the metrics above need no plotting stack."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10, 4))
    t, on = rec["t"], rec["lamp"]
    if any(_ok(v) for v in rec["raw_body"]):
        ax.plot(t, rec["raw_body"], label="body (raw, pre-gate)",
                color="#7f8c8d", linewidth=1.0, linestyle=":")
    ax.plot(t, rec["body"], label="body", color="#2471a3", linewidth=1.2)
    ax.plot(t, rec["ambient"], label="ambient", color="#c0392b", linewidth=1.2)
    for key, colour in (("body_sp", "#2471a3"), ("amb_sp", "#c0392b")):
        if any(_ok(s) for s in rec[key]):
            ax.plot(t, rec[key], color=colour, linestyle="--", linewidth=0.9)
    # actuator ON as shading, so duty is readable without a second axis
    for i in range(len(on)):
        if on[i] and _ok(t[i]):
            t_end = t[i + 1] if i + 1 < len(t) and _ok(t[i + 1]) else t[i]
            ax.axvspan(t[i], t_end, color="#f39c12", alpha=0.18, linewidth=0)
    ax.set_xlabel("time (s)")
    ax.set_ylabel("°C")
    ax.set_title(os.path.basename(rec["path"]))
    ax.legend(loc="best", fontsize=8)
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def _expand(paths: list[str]) -> list[str]:
    out: list[str] = []
    for p in paths:
        if os.path.isdir(p):
            out.extend(sorted(glob.glob(os.path.join(p, "**", "*.jsonl"), recursive=True)))
        else:
            out.append(p)
    return out


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Summarize session/recording .jsonl files (read-only).")
    ap.add_argument("paths", nargs="+", help="recording .jsonl files, or directories of them")
    ap.add_argument("--csv", default=None, help="write one row per session here")
    ap.add_argument("--xlsx", default=None,
                    help="write an Excel workbook: README, Summary, one sheet per session (needs openpyxl)")
    ap.add_argument("--plots", default=None, metavar="DIR",
                    help="write one overview PNG per session into DIR (needs matplotlib)")
    args = ap.parse_args(argv)

    files = _expand(args.paths)
    if not files:
        print("no .jsonl files found", file=sys.stderr)
        return 2

    rows, recs = [], []
    for path in files:
        rec = parse(path)
        if not rec["lamp"]:
            print(f"{path}: no samples — skipped", file=sys.stderr)
            continue
        rows.append(summarize(rec))
        recs.append(rec)
        print(f"== {path}")
        print(control_stats_text(rec))
        print()
        if args.plots:
            os.makedirs(args.plots, exist_ok=True)
            name = os.path.splitext(os.path.basename(path))[0] + ".png"
            plot(rec, os.path.join(args.plots, name))

    if args.xlsx and recs:
        write_xlsx(recs, args.xlsx)
        print(f"wrote {args.xlsx} ({len(recs)} session sheet(s))")
    if args.csv:
        write_csv(rows, args.csv)
        print(f"wrote {args.csv} ({len(rows)} session(s))")
    return 0 if rows else 1


if __name__ == "__main__":
    raise SystemExit(main())
